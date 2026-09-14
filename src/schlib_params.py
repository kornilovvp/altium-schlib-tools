#!/usr/bin/env python3
"""schlib_params.py -- parameter table (Excel / CSV / JSON) for Altium .SchLib files.

Usage:
  python schlib_params.py list  <lib.SchLib>
  python schlib_params.py show  <lib.SchLib> <LibRef>
  python schlib_params.py dump  <lib.SchLib> [-o table.xlsx|.csv|.json] [--sort alpha|frequency]
  python schlib_params.py diff  <lib.SchLib> <table.xlsx|.csv>
  python schlib_params.py apply <lib.SchLib> <table.xlsx|.csv> [-o out.SchLib] [--no-backup]

Table rules (dump / apply):
  * one row per component, key column "LibReference" (read-only)
  * fixed columns Description, DesignItemId, Designator (editable), Footprints (read-only)
  * one column per parameter name; a hatched cell means the component has no such parameter
  * apply: empty cell = no change; "<del>" removes the parameter; "<empty>" sets an empty value
    (creating the parameter when needed); any other text sets the value (creating when needed)
"""
import argparse
import csv
import json
import os
import sys
from typing import Dict, List, Optional, Tuple

import altium_schlib as asl

FIXED_COLUMNS = ["LibReference", "Description", "DesignItemId", "Designator", "Footprints"]
READ_ONLY_COLUMNS = {"LibReference", "Footprints"}
FIXED_ATTRS = {"Description": "description", "DesignItemId": "design_item_id", "Designator": "designator"}
DEL_TOKEN = "<del>"
EMPTY_TOKEN = "<empty>"
SHEET_NAME = "Parameters"

Row = Dict[str, Optional[str]]


class Change:
    """One planned edit.  kind: set | add | delete | attr"""

    def __init__(self, comp: asl.Component, kind: str, name: str, old: Optional[str], new: Optional[str]):
        self.comp, self.kind, self.name, self.old, self.new = comp, kind, name, old, new

    def __str__(self):
        ref = self.comp.lib_reference
        if self.kind == "add":
            return "%s: + %s = %r" % (ref, self.name, self.new)
        if self.kind == "delete":
            return "%s: - %s (was %r)" % (ref, self.name, self.old)
        return "%s: %s: %r -> %r" % (ref, self.name, self.old, self.new)


# ----------------------------------------------------------------------------------------------
# table <-> library
# ----------------------------------------------------------------------------------------------

def build_table(lib: asl.SchLib, sort: str = "frequency") -> Tuple[List[str], List[Row]]:
    """Rows are dicts column -> value; a missing parameter is represented by None."""
    columns = FIXED_COLUMNS + lib.all_param_names(sort)
    by_lower = {c.lower(): c for c in columns}
    rows: List[Row] = []
    for c in lib.components:
        row: Row = {"LibReference": c.lib_reference, "Description": c.description,
                    "DesignItemId": c.design_item_id, "Designator": c.designator,
                    "Footprints": "; ".join(c.footprints)}
        for p in c.params:
            col = by_lower.get(p.name.lower(), p.name)
            row.setdefault(col, p.value)
        rows.append(row)
    return columns, rows


def plan_changes(lib: asl.SchLib, columns: List[str], rows: List[Row]) -> Tuple[List[Change], List[str]]:
    changes: List[Change] = []
    warnings: List[str] = []
    for row in rows:
        ref = (row.get("LibReference") or "").strip()
        if not ref:
            continue
        comp = lib.component(ref)
        if comp is None:
            warnings.append("component %r not found in library, row skipped" % ref)
            continue
        for col in columns:
            if not col or col in READ_ONLY_COLUMNS:
                continue
            cell = row.get(col)
            if cell is None:
                continue
            cell = str(cell)
            if cell.strip() == "":
                continue                                   # empty cell = no change
            if col in FIXED_ATTRS:
                new = "" if cell == EMPTY_TOKEN else cell
                old = getattr(comp, FIXED_ATTRS[col])
                if new != old:
                    changes.append(Change(comp, "attr", col, old, new))
                continue
            p = comp.find_param(col)
            if cell == DEL_TOKEN:
                if p is not None:
                    changes.append(Change(comp, "delete", col, p.value, None))
                continue
            new = "" if cell == EMPTY_TOKEN else cell
            if p is None:
                changes.append(Change(comp, "add", col, None, new))
            elif p.value != new:
                changes.append(Change(comp, "set", col, p.value, new))
    return changes, warnings


def apply_changes(lib: asl.SchLib, changes: List[Change]) -> int:
    for ch in changes:
        if ch.kind == "delete":
            ch.comp.delete_param(ch.name)
        elif ch.kind in ("set", "add"):
            ch.comp.set_param(ch.name, ch.new)
        elif ch.kind == "attr":
            setattr(ch.comp, FIXED_ATTRS[ch.name], ch.new)
    return len(changes)


# ----------------------------------------------------------------------------------------------
# file formats
# ----------------------------------------------------------------------------------------------

def write_xlsx(path: str, columns: List[str], rows: List[Row]):
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.utils import get_column_letter

    wb = Workbook()
    ws = wb.active
    ws.title = SHEET_NAME
    missing_fill = PatternFill(fill_type="lightUp", fgColor="FF9E9E9E", bgColor="FFFFFFFF")
    ro_fill = PatternFill(fill_type="solid", fgColor="FFE7E6E6")
    head_font = Font(bold=True)
    for ci, col in enumerate(columns, start=1):
        cell = ws.cell(row=1, column=ci, value=col)
        cell.font = head_font
        cell.alignment = Alignment(horizontal="center")
    for r, row in enumerate(rows, start=2):
        for ci, col in enumerate(columns, start=1):
            cell = ws.cell(row=r, column=ci)
            cell.number_format = "@"                      # keep everything as text in Excel
            v = row.get(col)
            if v is None:
                if col not in FIXED_COLUMNS:
                    cell.fill = missing_fill
            else:
                cell.value = v
            if col in READ_ONLY_COLUMNS:
                cell.fill = ro_fill
    for ci, col in enumerate(columns, start=1):
        width = max([len(col)] + [len(str(row.get(col) or "")) for row in rows])
        ws.column_dimensions[get_column_letter(ci)].width = min(max(10, width + 2), 60)
    ws.freeze_panes = "B2"
    ws.auto_filter.ref = "A1:%s%d" % (get_column_letter(len(columns)), len(rows) + 1)
    hs = wb.create_sheet("Help")
    for line in [
        "One row per component.  Key column: LibReference (do not edit).",
        "Hatched cell = the component does not have that parameter.",
        "Empty cell = no change on apply.",
        "%s = remove the parameter from the component." % DEL_TOKEN,
        "%s = set an empty value (creates the parameter if needed)." % EMPTY_TOKEN,
        "Any other text = set the value (creates the parameter if needed).",
        "Add a new column with a parameter name to create a new parameter.",
        "Footprints is informational only.",
        "Apply with:  python schlib_params.py apply <lib.SchLib> <this file> -o <out.SchLib>",
    ]:
        hs.append([line])
    hs.column_dimensions["A"].width = 90
    wb.save(path)


def write_csv(path: str, columns: List[str], rows: List[Row]):
    with open(path, "w", newline="", encoding="utf-8-sig") as fh:
        w = csv.writer(fh, delimiter=";")
        w.writerow(columns)
        for row in rows:
            w.writerow(["" if row.get(c) is None else row.get(c) for c in columns])


def write_json(path: str, lib: asl.SchLib):
    data = {"file": lib.path, "warnings": lib.warnings, "components": []}
    for c in lib.components:
        data["components"].append({
            "LibReference": c.lib_reference, "Description": c.description, "DesignItemId": c.design_item_id,
            "Designator": c.designator, "PartCount": c.part_count, "storage": c.storage_name,
            "parameters": [{"name": p.name, "value": p.value, "hidden": p.hidden, "uid": p.unique_id} for p in c.params],
            "models": c.models, "pins": c.pins,
            "records": [{"index": i, "type": r.type_name, "fields": r.describe()} for i, r in enumerate(c.records)],
        })
    with open(path, "w", encoding="utf-8") as fh:
        json.dump(data, fh, ensure_ascii=False, indent=1, default=str)


def write_table(path: str, lib: asl.SchLib, sort: str = "frequency"):
    ext = os.path.splitext(path)[1].lower()
    columns, rows = build_table(lib, sort)
    if ext == ".xlsx":
        write_xlsx(path, columns, rows)
    elif ext == ".csv":
        write_csv(path, columns, rows)
    elif ext == ".json":
        write_json(path, lib)
    else:
        raise ValueError("unsupported output format %r (use .xlsx, .csv or .json)" % ext)


def _cell_to_str(v) -> Optional[str]:
    if v is None:
        return None
    if isinstance(v, float) and v.is_integer():
        v = int(v)
    return str(v)


def read_table(path: str) -> Tuple[List[str], List[Row]]:
    ext = os.path.splitext(path)[1].lower()
    rows: List[Row] = []
    if ext in (".xlsx", ".xlsm"):
        from openpyxl import load_workbook
        wb = load_workbook(path, data_only=True, read_only=True)
        ws = wb[SHEET_NAME] if SHEET_NAME in wb.sheetnames else wb.active
        it = ws.iter_rows(values_only=True)
        header = next(it, None) or ()
        columns = ["" if h is None else str(h).strip() for h in header]
        for vals in it:
            row: Row = {}
            for col, v in zip(columns, vals):
                if col and v is not None:
                    row[col] = _cell_to_str(v)
            if row.get("LibReference"):
                rows.append(row)
        wb.close()
    elif ext == ".csv":
        with open(path, "r", newline="", encoding="utf-8-sig") as fh:
            sample = fh.read(4096)
            fh.seek(0)
            try:
                dialect = csv.Sniffer().sniff(sample, delimiters=";,\t")
            except csv.Error:
                dialect = csv.excel
            reader = csv.reader(fh, dialect)
            header = next(reader, [])
            columns = [h.strip() for h in header]
            for vals in reader:
                row = {c: v for c, v in zip(columns, vals) if c and v != ""}
                if row.get("LibReference"):
                    rows.append(row)
    else:
        raise ValueError("unsupported table format %r (use .xlsx or .csv)" % ext)
    return [c for c in columns if c], rows


# ----------------------------------------------------------------------------------------------
# CLI
# ----------------------------------------------------------------------------------------------

def _load(path: str) -> asl.SchLib:
    lib = asl.SchLib.load(path)
    for w in lib.warnings:
        print("warning:", w, file=sys.stderr)
    return lib


def cmd_list(args):
    lib = _load(args.lib)
    counts = lib.param_name_counts()
    print("%s: %d components, %d parameter names" % (lib.path, len(lib.components), len(counts)))
    for name, n in sorted(counts.items(), key=lambda kv: (-kv[1], kv[0].lower())):
        print("  %3d/%d  %s" % (n, len(lib.components), name))
    print()
    for c in lib.components:
        print("%-30s %3d params %3d pins  %-40s  %s" % (c.lib_reference, len(c.params), c.pin_count,
                                                        c.description[:40], "; ".join(c.footprints)))


def cmd_show(args):
    lib = _load(args.lib)
    comp = lib.component(args.libref)
    if comp is None:
        sys.exit("component %r not found" % args.libref)
    print(comp.describe())


def cmd_dump(args):
    lib = _load(args.lib)
    if args.output:
        write_table(args.output, lib, args.sort)
        print("written", args.output)
        return
    columns, rows = build_table(lib, args.sort)
    print("\t".join(columns))
    for row in rows:
        print("\t".join("<none>" if row.get(c) is None else row.get(c) for c in columns))


def cmd_diff(args):
    lib = _load(args.lib)
    columns, rows = read_table(args.table)
    changes, warnings = plan_changes(lib, columns, rows)
    for w in warnings:
        print("warning:", w, file=sys.stderr)
    for ch in changes:
        print(ch)
    print("%d change(s)" % len(changes))
    return lib, changes


def cmd_apply(args):
    lib, changes = cmd_diff(args)
    if not changes:
        print("nothing to do")
        return
    apply_changes(lib, changes)
    out = lib.save(args.output, backup=not args.no_backup)
    print("saved", out)


def main(argv=None):
    ap = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    sub = ap.add_subparsers(dest="cmd", required=True)
    p = sub.add_parser("list", help="list components and parameter names")
    p.add_argument("lib")
    p.set_defaults(fn=cmd_list)
    p = sub.add_parser("show", help="dump every record of one component")
    p.add_argument("lib")
    p.add_argument("libref")
    p.set_defaults(fn=cmd_show)
    p = sub.add_parser("dump", help="export the parameter table (.xlsx / .csv / .json)")
    p.add_argument("lib")
    p.add_argument("-o", "--output")
    p.add_argument("--sort", choices=["frequency", "alpha"], default="frequency")
    p.set_defaults(fn=cmd_dump)
    p = sub.add_parser("diff", help="show what apply would change")
    p.add_argument("lib")
    p.add_argument("table")
    p.set_defaults(fn=cmd_diff)
    p = sub.add_parser("apply", help="apply a table to the library")
    p.add_argument("lib")
    p.add_argument("table")
    p.add_argument("-o", "--output", help="output .SchLib (default: overwrite input, keeping a .bak)")
    p.add_argument("--no-backup", action="store_true")
    p.set_defaults(fn=cmd_apply)
    args = ap.parse_args(argv)
    try:
        args.fn(args)
    except (asl.SchLibError, ValueError, OSError) as exc:
        sys.exit("error: %s" % exc)


if __name__ == "__main__":
    if hasattr(sys.stdout, "reconfigure"):
        sys.stdout.reconfigure(encoding="utf-8", errors="replace")
    main()
